import json
import sys
import tempfile
import threading
import unittest
import urllib.request
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'pc'))
from core import Store, Server, HEADERS
from openpyxl import Workbook, load_workbook

class ReceiverTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.store=Store(self.temp.name)
        self.row={'id':'r1','employee_id':'e1','name':'김직원','time':'2026-10-08T09:00:00+09:00','kind':'출근','revision':1,'reason':''}
    def test_retry_and_revision(self):
        self.store.receive([self.row]);self.store.receive([self.row]);self.assertEqual(len(self.store.rows()),1)
        changed=dict(self.row,time='2026-10-08T08:55:00+09:00',revision=2,reason='시간 확인')
        self.store.receive([changed]);self.store.receive([self.row]);self.assertEqual(self.store.rows()[0]['revision'],2)
        with self.store.connect() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM audit').fetchone()[0],1)
    def test_batch_atomicity_and_conflict(self):
        self.store.receive([self.row]);bad=dict(self.row,name='다른 사람')
        with self.assertRaises(ValueError):self.store.receive([dict(self.row,id='r2'),bad])
        self.assertEqual(len(self.store.rows()),1)
        with self.assertRaises(ValueError):self.store.receive([dict(self.row,kind='외출')])
    def test_excel_and_formula_safety(self):
        self.store.receive([dict(self.row,name='=1+1')]);path=self.store.export()[0];ws=load_workbook(path).active
        self.assertEqual([c.value for c in ws[1]],HEADERS);self.assertEqual(ws['A2'].data_type,'s');self.assertEqual(ws['B2'].value.hour,9)
        self.assertEqual(ws.max_column,3)
    def test_locked_excel_preserves_database(self):
        self.store.receive([self.row]);self.store.export()
        with patch('core.os.replace',side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):self.store.export()
        self.assertEqual(len(self.store.rows()),1);self.store.export()
    def test_request_not_lost_during_upload(self):
        self.store.request();seq=self.store.control()['requested'];self.store.request();self.store.complete(seq)
        self.assertEqual(self.store.control()['completed'],1);self.assertEqual(self.store.control()['requested'],2)
        reopened=Store(self.temp.name);self.assertEqual(reopened.control()['requested'],2)
    def test_template_mapping(self):
        self.store.receive([self.row]);src=Path(self.temp.name)/'template.xlsx';dest=Path(self.temp.name)/'out.xlsx'
        wb=Workbook();wb.active.title='양식';wb.active['B1']='유지';wb.save(src)
        self.store.export_template(src,dest,'양식',5,[2,4,6],'2026-10');ws=load_workbook(dest).active
        self.assertEqual(ws['B1'].value,'유지');self.assertEqual(ws['B5'].value,'김직원');self.assertEqual(ws['F5'].value,'출근');self.assertIsNone(load_workbook(src).active['B5'].value)
    def test_http_auth_and_receipt(self):
        s=Server(('127.0.0.1',0),self.store,'x'*32);threading.Thread(target=s.serve_forever,daemon=True).start()
        self.addCleanup(s.server_close);self.addCleanup(s.shutdown)
        base='http://127.0.0.1:'+str(s.server_port)
        opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with self.assertRaises(urllib.error.HTTPError) as e:opener.open(base+'/control')
        self.assertEqual(e.exception.code,401)
        e.exception.close()
        req=urllib.request.Request(base+'/records',json.dumps({'records':[self.row]}).encode(),{'Authorization':'Bearer '+'x'*32,'Content-Type':'application/json'})
        with opener.open(req) as r:self.assertEqual(json.load(r)['accepted'],1)
        self.assertEqual(len(self.store.rows()),1)

if __name__=='__main__':unittest.main()
