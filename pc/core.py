"""Durable, idempotent receiver and atomic Excel snapshots. No attendance data on GitHub."""
from __future__ import annotations
from contextlib import contextmanager, closing
import secrets
import time
import hmac
import ipaddress
import json
import os
import re
import sqlite3
import tempfile
import threading
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from openpyxl import Workbook, load_workbook

KST = timezone(timedelta(hours=9))
HEADERS = ['직원이름', '시간', '출근인지 퇴근인지']

class Store:
    def __init__(self, folder):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.db = self.folder / 'attendance.db'
        self.export_lock = threading.Lock()
        self.export_error = ''
        with self.connect() as c:
            c.executescript('''
            CREATE TABLE IF NOT EXISTS records(id TEXT PRIMARY KEY, employee_id TEXT NOT NULL,
              name TEXT NOT NULL, time TEXT NOT NULL, kind TEXT NOT NULL, revision INTEGER NOT NULL,
              reason TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, record_id TEXT, old_json TEXT,
              new_json TEXT, received_at TEXT);
            CREATE TABLE IF NOT EXISTS control(id INTEGER PRIMARY KEY CHECK(id=1), requested INTEGER,
              completed INTEGER);
            INSERT OR IGNORE INTO control VALUES(1,0,0);
            ''')
    @contextmanager
    def connect(self):
        c = sqlite3.connect(self.db, timeout=15)
        c.row_factory = sqlite3.Row
        try:
            with c:
                yield c
        finally:
            c.close()
    @staticmethod
    def validate(r):
        if not isinstance(r, dict): raise ValueError('record must be an object')
        for k in ('id', 'employee_id', 'name', 'time', 'kind', 'reason'):
            if not isinstance(r.get(k), str): raise ValueError('invalid ' + k)
        if not r['id'] or len(r['id']) > 80 or not r['employee_id'] or len(r['employee_id']) > 80: raise ValueError('invalid id')
        if not r['name'].strip() or len(r['name']) > 40 or len(r['reason']) > 1000: raise ValueError('invalid name/reason')
        if r['kind'] not in ('출근', '퇴근'): raise ValueError('invalid kind')
        if type(r.get('revision')) is not int or not 1 <= r['revision'] <= 1000000: raise ValueError('invalid revision')
        when = datetime.fromisoformat(r['time'])
        if when.tzinfo is None or when.utcoffset() != timedelta(hours=9): raise ValueError('time must be KST with offset')
        if not 2020 <= when.year <= 2100: raise ValueError('invalid year')
        # Reject XML-illegal controls before they reach an Excel cell.
        if any(ord(ch) < 32 and ch not in '\t\n\r' for ch in r['name'] + r['reason']): raise ValueError('invalid control character')
        return {k: r[k] for k in ('id', 'employee_id', 'name', 'time', 'kind', 'revision', 'reason')}
    def receive(self, records):
        if not isinstance(records, list) or len(records) > 200: raise ValueError('batch limit 200')
        rows = [self.validate(r) for r in records]
        with self.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            for r in rows:
                old = c.execute('SELECT * FROM records WHERE id=?', (r['id'],)).fetchone()
                if old:
                    if old['employee_id'] != r['employee_id']: raise ValueError('employee identity changed')
                    if old['revision'] > r['revision']: continue
                    if old['revision'] == r['revision']:
                        if dict(old) != r: raise ValueError('same revision has different contents')
                        continue
                    c.execute('INSERT INTO audit(record_id,old_json,new_json,received_at) VALUES(?,?,?,?)',
                              (r['id'], json.dumps(dict(old), ensure_ascii=False), json.dumps(r, ensure_ascii=False), datetime.now(KST).isoformat()))
                c.execute('''INSERT INTO records VALUES(:id,:employee_id,:name,:time,:kind,:revision,:reason)
                 ON CONFLICT(id) DO UPDATE SET name=excluded.name,time=excluded.time,kind=excluded.kind,
                 revision=excluded.revision,reason=excluded.reason''', r)
        # ACK concerns the committed database, independent of Excel file locks.
        return len(rows)
    def control(self):
        with self.connect() as c:
            r = dict(c.execute('SELECT requested,completed FROM control WHERE id=1').fetchone())
            r['count'] = c.execute('SELECT COUNT(*) FROM records').fetchone()[0]
            return r
    def request(self):
        with self.connect() as c: c.execute('UPDATE control SET requested=requested+1 WHERE id=1')
    def complete(self, sequence):
        if type(sequence) is not int or sequence < 0: raise ValueError('invalid sequence')
        with self.connect() as c:
            c.execute('UPDATE control SET completed=MAX(completed,MIN(?,requested)) WHERE id=1', (sequence,))
    def rows(self, month=None):
        if month is not None and not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month): raise ValueError('month must be YYYY-MM')
        with self.connect() as c:
            return [dict(r) for r in c.execute('SELECT * FROM records WHERE time LIKE ? ORDER BY time,name,id', ((month or '') + '%',))]
    @staticmethod
    def cell_row(sheet, row, values, columns=(1,2,3)):
        for col, value in zip(columns, values):
            cell = sheet.cell(row, col, value)
            if isinstance(value, str): cell.data_type = 's'  # Names beginning with = are literal text.
            if isinstance(value, datetime): cell.number_format = 'yyyy-mm-dd hh:mm:ss'
    @staticmethod
    def values(r): return (r['name'], datetime.fromisoformat(r['time']).astimezone(KST).replace(tzinfo=None), r['kind'])
    @staticmethod
    def save_atomic(wb, path):
        path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix='.attendance-', suffix='.xlsx', dir=path.parent)
        os.close(fd)
        try: wb.save(name); os.replace(name, path)
        finally:
            if os.path.exists(name): os.unlink(name)
    def export(self, folder=None):
        destination = Path(folder) if folder is not None else self.folder / '엑셀'
        with self.export_lock:
            rows = self.rows(); months = sorted({r['time'][:7] for r in rows})
            outputs = []
            for month in months:
                wb=Workbook(); ws=wb.active; ws.title='근태기록'
                self.cell_row(ws, 1, HEADERS)
                for idx,r in enumerate((r for r in rows if r['time'].startswith(month)),2): self.cell_row(ws,idx,self.values(r))
                ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
                for col,width in [('A',24),('B',25),('C',24)]: ws.column_dimensions[col].width=width
                path=destination/f'근태기록_{month}.xlsx'; self.save_atomic(wb,path); outputs.append(path)
            self.export_error=''
            return outputs
    def export_template(self, template, destination, sheet, start_row, columns, month):
        """Map a flat table to an existing xlsx; source is never overwritten."""
        template=Path(template); destination=Path(destination)
        if template.suffix.lower() != '.xlsx' or destination.suffix.lower() != '.xlsx': raise ValueError('xlsx only')
        if template.resolve() == destination.resolve(): raise ValueError('원본 양식과 출력 파일은 달라야 합니다.')
        if start_row < 1 or len(columns)!=3 or len(set(columns))!=3 or any(not 1<=c<=16384 for c in columns): raise ValueError('잘못된 셀 위치')
        wb=load_workbook(template)
        if sheet not in wb.sheetnames: raise ValueError('시트 이름을 확인하세요.')
        ws=wb[sheet]; rows=self.rows(month)
        for idx,r in enumerate(rows,start_row): self.cell_row(ws,idx,self.values(r),columns)
        self.save_atomic(wb,destination)
        return len(rows)
    def backup(self, destination):
        with self.connect() as src, closing(sqlite3.connect(destination)) as dst: src.backup(dst)

class Pairing:
    """A local administrator opens a five-minute, single-use pairing window."""
    def __init__(self, clock=time.monotonic):
        self.clock=clock; self.lock=threading.Lock(); self.code=None; self.expires=0; self.attempts=0
    def issue(self):
        with self.lock:
            self.code=f'{secrets.randbelow(1000000):06d}'; self.expires=self.clock()+300; self.attempts=0
            return self.code
    def consume(self, code):
        with self.lock:
            if self.code is None or self.clock()>=self.expires or self.attempts>=5: return False
            self.attempts+=1
            if not isinstance(code,str) or not hmac.compare_digest(code.encode(),self.code.encode()): return False
            self.code=None
            return True
    def status(self):
        with self.lock:
            if self.code is None: return '연결 완료 또는 코드 미발급'
            if self.clock()>=self.expires or self.attempts>=5: return '코드 만료 · 새 코드를 발급하세요'
            return f'{self.code}  ·  {int(self.expires-self.clock())}초 남음'

class Server(ThreadingHTTPServer):
    daemon_threads=True
    def __init__(self, address, store, token, pairing=None):
        self.store=store; self.token=token; self.pairing=pairing or Pairing()
        super().__init__(address, Handler)

class Handler(BaseHTTPRequestHandler):
    def setup(self):
        super().setup(); self.connection.settimeout(15)
    def log_message(self, *_): pass  # Never log bearer credentials or employee data.
    def answer(self, status, value):
        b=json.dumps(value, ensure_ascii=False).encode('utf-8'); self.send_response(status)
        self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b)
    def authorized(self):
        ok=hmac.compare_digest(self.headers.get('Authorization','').encode('utf-8'), ('Bearer '+self.server.token).encode('utf-8'))
        if not ok: self.answer(401, {'error':'unauthorized'})
        return ok
    def do_GET(self):
        if not self.authorized(): return
        if self.path!='/control': self.answer(404, {'error':'not found'}); return
        self.answer(200,self.server.store.control())
    def do_POST(self):
        if self.path != '/pair' and not self.authorized(): return
        try:
            n=int(self.headers.get('Content-Length','0'))
            if not 0<n<=1000000: raise ValueError('invalid size')
            body=json.loads(self.rfile.read(n))
            if self.path=='/pair':
                if not self.server.pairing.consume(body.get('code')):
                    self.answer(403, {'error':'pairing code invalid or expired'}); return
                result={'token':self.server.token}
            elif self.path=='/records': result={'accepted':self.server.store.receive(body['records'])}
            elif self.path=='/complete': self.server.store.complete(body['sequence']); result={'ok':True}
            else: self.answer(404, {'error':'not found'}); return
            self.answer(200,result)
        except (ValueError,KeyError,TypeError): self.answer(400, {'error':'invalid request'})
        except Exception: self.answer(500, {'error':'storage failure; retry'})

def tailscale_ip(value):
    try: return ipaddress.ip_address(value) in ipaddress.ip_network('100.64.0.0/10')
    except ValueError: return False
